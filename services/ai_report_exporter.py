from __future__ import annotations

import html
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from services.report_fonts import html_font_css, set_docx_fonts


def markdown_to_html_document(markdown_text: str, title: str = "AI 飞行日志分析报告") -> str:
    body = []
    for line in markdown_text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("# "):
            body.append(f"<h1>{html.escape(stripped[2:])}</h1>")
        elif stripped.startswith("## "):
            body.append(f"<h2>{html.escape(stripped[3:])}</h2>")
        elif stripped.startswith("### "):
            body.append(f"<h3>{html.escape(stripped[4:])}</h3>")
        elif stripped.startswith("- ") or stripped.startswith("* "):
            body.append(f"<p>• {html.escape(stripped[2:])}</p>")
        elif stripped.startswith("| "):
            body.append(f"<pre>{html.escape(stripped)}</pre>")
        else:
            body.append(f"<p>{html.escape(stripped)}</p>")
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <title>{html.escape(title)}</title>
  <style>
  {html_font_css()}
  body {{ max-width: 980px; margin: 28px auto; padding: 0 24px; color: #172027; line-height: 1.72; }}
  h1 {{ color: #0b2545; border-bottom: 2px solid #18b7c8; padding-bottom: 10px; }}
  h2 {{ color: #0f6b7a; margin-top: 28px; }}
  h3 {{ color: #243946; }}
  pre {{ white-space: pre-wrap; background: #f4f8fa; border: 1px solid #d9e6eb; padding: 8px; }}
  .notice {{ margin-top: 24px; padding: 12px; background: #eef9fb; border-left: 4px solid #18b7c8; }}
  </style>
</head>
<body>
{chr(10).join(body)}
<div class="notice">AI 报告基于算法分析结果生成，仅用于辅助分析，不能替代人工工程判断。</div>
</body>
</html>"""


def markdown_to_docx_document(markdown_text: str, output_path: Path) -> None:
    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("AI 飞行日志分析报告")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(11, 37, 69)

    for line in markdown_text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        else:
            doc.add_paragraph(stripped)

    note = doc.add_paragraph()
    note.add_run("安全边界：").bold = True
    note.add_run("AI 报告基于算法分析结果生成，仅用于辅助分析，不能替代人工工程判断；AI 不直接控制飞控，不在飞行中自动修改 PID。")
    set_docx_fonts(doc)
    doc.save(output_path)


def export_ai_report_files(report: dict[str, Any], output_dir: Path, stem: str) -> dict[str, str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    markdown_text = report["reportMarkdown"]
    md_path = output_dir / f"{stem}.md"
    html_path = output_dir / f"{stem}.html"
    docx_path = output_dir / f"{stem}.docx"
    md_path.write_text(markdown_text, encoding="utf-8")
    html_path.write_text(markdown_to_html_document(markdown_text), encoding="utf-8")
    markdown_to_docx_document(markdown_text, docx_path)
    return {
        "markdownPath": str(md_path),
        "htmlPath": str(html_path),
        "docxPath": str(docx_path),
        "markdownUrl": f"/reports/{md_path.name}",
        "htmlUrl": f"/reports/{html_path.name}",
        "docxUrl": f"/reports/{docx_path.name}",
    }
